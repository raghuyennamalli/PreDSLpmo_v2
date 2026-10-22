#!/usr/bin/env python3

import os
import re
import random
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

BASE_DIR = "/home/ragothaman/Documents/Sweety/PredLPMO/New_Version/Cluster_Aware_Spliting"

POSITIVE_DIR = os.path.join(BASE_DIR, "Positive")
NEGATIVE_DIR = os.path.join(BASE_DIR, "Negative")

OUTPUT_DIR = os.path.join(BASE_DIR, "Cluster_Split_Output")

TRAIN_DIR = os.path.join(OUTPUT_DIR, "Training")
VAL_DIR = os.path.join(OUTPUT_DIR, "Validation")

EXCEL_OUTPUT = os.path.join(OUTPUT_DIR, "cluster_split_report.xlsx")

TRAIN_RATIO = 0.60
RANDOM_SEED = 42
POSITIVE_FAMILIES = [
    "AA9",
    "AA10",
    "AA11",
    "AA13",
    "AA14",
    "AA15",
    "AA16",
    "AA17"
]

NEGATIVE_FAMILIES = [
    "AA1",
    "AA2",
    "AA3",
    "AA4",
    "AA5",
    "AA6",
    "AA7",
    "AA8",
    "AA12"
]


os.makedirs(TRAIN_DIR, exist_ok=True)
os.makedirs(VAL_DIR, exist_ok=True)


def read_fasta(path):
    """
    Read FASTA.

    Returns:
        dict:
            sequence_id -> complete FASTA record
    """

    sequences = {}

    current_id = None
    current_header = None
    current_sequence = []

    with open(path, "r") as f:

        for line in f:

            line = line.rstrip("\n")

            if not line:
                continue

            if line.startswith(">"):

                # Save previous record
                if current_id is not None:

                    sequences[current_id] = {
                        "header": current_header,
                        "sequence": "".join(current_sequence)
                    }

                current_header = line[1:].strip()

                current_id = current_header.split()[0]

                current_sequence = []

            else:

                current_sequence.append(line.strip())

        if current_id is not None:

            sequences[current_id] = {
                "header": current_header,
                "sequence": "".join(current_sequence)
            }

    return sequences

def normalize_id(identifier):
    """
    Normalize an identifier for comparison.

    IMPORTANT:
    Do NOT split JGI identifiers at '|'.
    """

    identifier = identifier.strip()

    if identifier.startswith(">"):
        identifier = identifier[1:]

    identifier = identifier.rstrip(".")

    identifier = identifier.split()[0]

    return identifier

def build_prefix_index(fasta_ids):

    index = defaultdict(list)

    for seq_id in fasta_ids:

        seq_id = normalize_id(seq_id)

        for length in range(8, len(seq_id) + 1):

            prefix = seq_id[:length]

            index[prefix].append(seq_id)

    return index


def match_cluster_id(raw_id, fasta_ids, prefix_index):

    raw_id = normalize_id(raw_id)


    if raw_id in fasta_ids:
        return raw_id


    candidates = prefix_index.get(raw_id, [])

    if len(candidates) == 1:
        return candidates[0]

    raw_prefix = raw_id.split("...")[0]

    candidates = []

    for fasta_id in fasta_ids:

        if fasta_id.startswith(raw_prefix):
            candidates.append(fasta_id)

    candidates = list(set(candidates))

    if len(candidates) == 1:
        return candidates[0]

    if len(candidates) == 0:

        raise RuntimeError(
            f"Could not match CD-HIT identifier to FASTA:\n"
            f"    {raw_id}"
        )

    else:

        raise RuntimeError(
            f"Ambiguous CD-HIT identifier:\n"
            f"    {raw_id}\n"
            f"Possible FASTA IDs:\n"
            + "\n".join(
                f"    {x}" for x in candidates[:20]
            )
        )


def parse_cd_hit_clusters(clstr_path, fasta_sequences):


    fasta_ids = set(fasta_sequences.keys())

    prefix_index = build_prefix_index(fasta_ids)

    clusters = {}

    current_cluster = None

    total_members = 0

    print("\nReading CD-HIT cluster file:")
    print(clstr_path)

    with open(clstr_path, "r") as f:

        for line in f:

            line = line.strip()

            if not line:
                continue

            if line.startswith(">Cluster"):

                cluster_number = int(
                    line.split()[1]
                )

                current_cluster = cluster_number

                clusters[current_cluster] = []

                continue


            if current_cluster is None:
                continue


            match = re.search(
                r">(.+?)(?:\.\.\.)?(?:\s+\*|\s+at\s+[\d.]+%|\s*$)",
                line
            )

            if not match:

                raise RuntimeError(
                    f"Could not parse cluster line:\n{line}"
                )

            raw_id = match.group(1)

            fasta_id = match_cluster_id(
                raw_id,
                fasta_ids,
                prefix_index
            )

            clusters[current_cluster].append(
                fasta_id
            )

            total_members += 1

    print("\nCD-HIT parsing summary")

    print(
        f"Number of clusters : {len(clusters)}"
    )

    print(
        f"Number of members  : {total_members}"
    )

    for cluster_id, members in clusters.items():

        if len(members) != len(set(members)):

            raise RuntimeError(
                f"Duplicate sequence ID detected in "
                f"cluster {cluster_id}"
            )
    clustered_ids = set()

    for members in clusters.values():
        clustered_ids.update(members)

    missing_from_clusters = (
        fasta_ids - clustered_ids
    )

    extra_in_clusters = (
        clustered_ids - fasta_ids
    )

    print(
        f"FASTA sequences     : {len(fasta_ids)}"
    )

    print(
        f"Clustered sequences  : {len(clustered_ids)}"
    )

    print(
        f"Missing from .clstr  : {len(missing_from_clusters)}"
    )

    print(
        f"Extra IDs            : {len(extra_in_clusters)}"
    )

    if missing_from_clusters:

        print("\nWARNING: FASTA sequences absent from .clstr")

        for seq_id in list(missing_from_clusters)[:20]:
            print("   ", seq_id)

    if extra_in_clusters:

        raise RuntimeError(
            "Some cluster IDs were not present in FASTA."
        )

    return clusters


def cluster_aware_split(
    clusters,
    train_ratio=0.60,
    seed=42
):

    """
    Assign ENTIRE clusters to either training or validation. No cluster can appear in both sets.
    """

    random.seed(seed)

    cluster_ids = list(clusters.keys())

    random.shuffle(cluster_ids)

    total_sequences = sum(
        len(clusters[c])
        for c in cluster_ids
    )

    target_training = (
        total_sequences * train_ratio
    )

    training_clusters = []
    validation_clusters = []

    training_sequences = 0

    for cluster_id in cluster_ids:

        cluster_size = len(
            clusters[cluster_id]
        )

        if (
            training_sequences + cluster_size
            <= target_training
        ):

            training_clusters.append(
                cluster_id
            )

            training_sequences += cluster_size

        else:

            validation_clusters.append(
                cluster_id
            )


    return (
        training_clusters,
        validation_clusters
    )



def write_fasta(
    sequence_ids,
    fasta_sequences,
    output_path
):

    with open(output_path, "w") as f:

        for seq_id in sequence_ids:

            record = fasta_sequences[seq_id]

            f.write(
                ">" + record["header"] + "\n"
            )

            sequence = record["sequence"]

            for i in range(0, len(sequence), 80):

                f.write(
                    sequence[i:i+80] + "\n"
                )


def process_family(
    family,
    input_dir,
    prefix,
    output_prefix
):

    print("\n")
    print("=" * 70)
    print(f"Processing: {prefix} {family}")
    print("=" * 70)

    fasta_path = os.path.join(
        input_dir,
        f"{family}_combined_cleaned.fasta"
    )

    clstr_path = os.path.join(
        input_dir,
        f"{family}_combined_cd_hit.fasta.clstr"
    )


    if not os.path.exists(fasta_path):

        raise FileNotFoundError(
            f"FASTA not found:\n{fasta_path}"
        )

    if not os.path.exists(clstr_path):

        raise FileNotFoundError(
            f"CLSTR not found:\n{clstr_path}"
        )


    fasta_sequences = read_fasta(
        fasta_path
    )

    print(
        f"FASTA sequences: {len(fasta_sequences)}"
    )


    clusters = parse_cd_hit_clusters(
        clstr_path,
        fasta_sequences
    )

    print(
        f"CD-HIT clusters: {len(clusters)}"
    )


    training_clusters, validation_clusters = (
        cluster_aware_split(
            clusters,
            TRAIN_RATIO,
            RANDOM_SEED
        )
    )


    training_ids = []

    for cluster_id in training_clusters:

        training_ids.extend(
            clusters[cluster_id]
        )

    validation_ids = []

    for cluster_id in validation_clusters:

        validation_ids.extend(
            clusters[cluster_id]
        )

    training_ids = list(
        dict.fromkeys(training_ids)
    )

    validation_ids = list(
        dict.fromkeys(validation_ids)
    )


    sequence_overlap = (
        set(training_ids)
        &
        set(validation_ids)
    )

    if sequence_overlap:

        raise RuntimeError(
            f"ERROR: Sequence overlap detected "
            f"in {family}: "
            f"{len(sequence_overlap)} sequences"
        )


    cluster_overlap = (
        set(training_clusters)
        &
        set(validation_clusters)
    )

    if cluster_overlap:

        raise RuntimeError(
            f"ERROR: Cluster overlap detected "
            f"in {family}: "
            f"{len(cluster_overlap)} clusters"
        )


    all_split_ids = (
        set(training_ids)
        |
        set(validation_ids)
    )

    fasta_ids = set(
        fasta_sequences.keys()
    )

    missing = fasta_ids - all_split_ids
    duplicated = (
        len(training_ids)
        + len(validation_ids)
        - len(all_split_ids)
    )

    if missing:

        print(
            f"WARNING: {len(missing)} FASTA sequences "
            f"were not assigned to either split."
        )

        for seq_id in list(missing)[:20]:
            print("   ", seq_id)

    if duplicated != 0:

        raise RuntimeError(
            f"Unexpected duplicate assignment: "
            f"{duplicated}"
        )

    total = len(fasta_sequences)

    train_percent = (
        len(training_ids) / total * 100
    )

    val_percent = (
        len(validation_ids) / total * 100
    )


    train_output = os.path.join(
        TRAIN_DIR,
        f"{output_prefix}_{family}.fasta"
    )

    val_output = os.path.join(
        VAL_DIR,
        f"{output_prefix}_{family}.fasta"
    )


    write_fasta(
        training_ids,
        fasta_sequences,
        train_output
    )

    write_fasta(
        validation_ids,
        fasta_sequences,
        val_output
    )


    print("\nRESULT")

    print(
        f"Total sequences      : {total}"
    )

    print(
        f"CD-HIT clusters      : {len(clusters)}"
    )

    print(
        f"Training sequences   : {len(training_ids)} "
        f"({train_percent:.2f}%)"
    )

    print(
        f"Validation sequences : {len(validation_ids)} "
        f"({val_percent:.2f}%)"
    )

    print(
        f"Training clusters    : {len(training_clusters)}"
    )

    print(
        f"Validation clusters  : {len(validation_clusters)}"
    )

    print(
        f"Cluster overlap      : {len(cluster_overlap)}"
    )

    print(
        f"Sequence overlap     : {len(sequence_overlap)}"
    )

    print(
        f"Training file        : {train_output}"
    )

    print(
        f"Validation file      : {val_output}"
    )

    result = {
        "family": family,
        "prefix": prefix,
        "total": total,
        "clusters": len(clusters),
        "training": len(training_ids),
        "validation": len(validation_ids),
        "training_clusters": len(training_clusters),
        "validation_clusters": len(validation_clusters),
        "cluster_overlap": len(cluster_overlap),
        "sequence_overlap": len(sequence_overlap)
    }

    cluster_rows = []

    training_cluster_set = set(training_clusters)

    for cluster_id, members in clusters.items():

        split_name = (
            "Training"
            if cluster_id in training_cluster_set
            else "Validation"
        )

        cluster_rows.append({
            "prefix": prefix,
            "family": family,
            "cluster_id": cluster_id,
            "split": split_name,
            "num_sequences": len(members),
            "sequence_ids": ", ".join(members)
        })

    return result, cluster_rows


def write_excel_summary(results, all_cluster_rows, output_path):
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, name="Arial", color="FFFFFF")
    header_fill = PatternFill(
        start_color="4472C4",
        end_color="4472C4",
        fill_type="solid"
    )
    body_font = Font(name="Arial")
    center = Alignment(horizontal="center")


    ws1 = wb.active
    ws1.title = "Final Summary"

    summary_headers = [
        "Class", "Family", "Total_sequences", "CDHIT_clusters",
        "Training_sequences", "Validation_sequences",
        "Training_clusters", "Validation_clusters",
        "Training_percent", "Validation_percent",
        "Cluster_overlap", "Sequence_overlap"
    ]

    ws1.append(summary_headers)

    for col in range(1, len(summary_headers) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for r in results:

        train_pct = round(r["training"] / r["total"] * 100, 2) if r["total"] else 0
        val_pct = round(r["validation"] / r["total"] * 100, 2) if r["total"] else 0

        ws1.append([
            r["prefix"],
            r["family"],
            r["total"],
            r["clusters"],
            r["training"],
            r["validation"],
            r["training_clusters"],
            r["validation_clusters"],
            train_pct,
            val_pct,
            r["cluster_overlap"],
            r["sequence_overlap"]
        ])

    for row in ws1.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font

    # Column widths
    widths1 = [10, 10, 16, 16, 20, 22, 18, 20, 18, 20, 16, 16]
    for i, w in enumerate(widths1, start=1):
        ws1.column_dimensions[
            openpyxl.utils.get_column_letter(i)
        ].width = w

    ws1.freeze_panes = "A2"


    ws2 = wb.create_sheet("Cluster Assignments")

    cluster_headers = [
        "Class", "Family", "Cluster_ID", "Split",
        "Num_Sequences", "Sequence_IDs"
    ]

    ws2.append(cluster_headers)

    for col in range(1, len(cluster_headers) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # Sort for readability: by class, family, then cluster ID
    all_cluster_rows_sorted = sorted(
        all_cluster_rows,
        key=lambda x: (x["prefix"], x["family"], x["cluster_id"])
    )

    for row in all_cluster_rows_sorted:

        ws2.append([
            row["prefix"],
            row["family"],
            row["cluster_id"],
            row["split"],
            row["num_sequences"],
            row["sequence_ids"]
        ])

    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font

    widths2 = [10, 10, 12, 12, 15, 60]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[
            openpyxl.utils.get_column_letter(i)
        ].width = w

    ws2.freeze_panes = "A2"

    wb.save(output_path)

    print(f"\nExcel report written to:\n{output_path}")


if __name__ == "__main__":

    print("=" * 70)
    print("CLUSTER-AWARE TRAINING / VALIDATION SPLITTING")
    print("=" * 70)

    print(
        f"\nTraining ratio    : {TRAIN_RATIO}"
    )

    print(
        f"Validation ratio  : {1 - TRAIN_RATIO}"
    )

    print(
        f"Random seed       : {RANDOM_SEED}"
    )

    results = []
    all_cluster_rows = []


    for family in POSITIVE_FAMILIES:

        result, cluster_rows = process_family(
            family=family,
            input_dir=POSITIVE_DIR,
            prefix="LPMO",
            output_prefix="LPMO"
        )

        results.append(result)
        all_cluster_rows.extend(cluster_rows)


    for family in NEGATIVE_FAMILIES:

        result, cluster_rows = process_family(
            family=family,
            input_dir=NEGATIVE_DIR,
            prefix="NON-LPMO",
            output_prefix="NON-LPMO"
        )

        results.append(result)
        all_cluster_rows.extend(cluster_rows)


    print("\n")
    print("=" * 70)
    print("FINAL SUMMARY")
    print("=" * 70)

    print(
        f"{'Family':<10}"
        f"{'Total':>10}"
        f"{'Train':>10}"
        f"{'Val':>10}"
        f"{'Clusters':>12}"
        f"{'Train Cl':>12}"
        f"{'Val Cl':>12}"
        f"{'Cl Overlap':>12}"
        f"{'Seq Overlap':>12}"
    )

    print("-" * 100)

    for r in results:

        print(
            f"{r['family']:<10}"
            f"{r['total']:>10}"
            f"{r['training']:>10}"
            f"{r['validation']:>10}"
            f"{r['clusters']:>12}"
            f"{r['training_clusters']:>12}"
            f"{r['validation_clusters']:>12}"
            f"{r['cluster_overlap']:>12}"
            f"{r['sequence_overlap']:>12}"
        )

    write_excel_summary(
        results,
        all_cluster_rows,
        EXCEL_OUTPUT
    )

    print("\nDONE.")
